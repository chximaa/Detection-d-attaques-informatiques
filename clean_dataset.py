"""
clean_dataset.py
----------------
Lit l'Excel brut (cyber_data_raw.xlsx), nettoie et prépare les données pour
l'entraînement d'un modèle de détection d'attaques, puis exporte un Excel
propre (cyber_data_clean.xlsx) avec les feuilles :
    train / test / label_mapping / scaler_params / cleaning_report
Chaque étape est affichée dans le terminal.

Utilisation : python clean_dataset.py
              python clean_dataset.py --input mon_fichier.xlsx --output sortie.xlsx

Remarque anti-fuite de données (data leakage) : le découpage train/test est fait
AVANT l'imputation, le clipping et la normalisation. Toutes les statistiques
(médianes, bornes, moyennes, écarts-types) sont calculées sur le train
uniquement puis appliquées au test.
"""
import argparse

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler

# ----------------------------------------------------------------------------
# Configuration (à adapter si vos colonnes changent)
# ----------------------------------------------------------------------------
TARGET = "label"
ID_COLS = ["flow_id", "src_ip", "dst_ip", "timestamp", "src_port"]   # supprimées avant entraînement
NUM_COLS = ["src_port", "dst_port", "duration", "src_bytes", "dst_bytes",
            "packets_fwd", "packets_bwd", "failed_logins", "conn_count"]
CAT_COLS = ["protocol", "service", "flag"]
PORT_COLS = ["src_port", "dst_port"]
NON_NEGATIVE = ["duration", "src_bytes", "dst_bytes", "packets_fwd",
                "packets_bwd", "failed_logins", "conn_count"]
LABEL_MAP = {"normal": "Normal", "dos": "DoS", "portscan": "PortScan", "bruteforce": "BruteForce"}
LABEL_ORDER = ["Normal", "BruteForce", "DoS", "PortScan"]            # Normal = 0
SKEWED = ["duration", "src_bytes", "dst_bytes", "packets_fwd", "packets_bwd",
          "failed_logins", "conn_count", "total_bytes", "total_packets",
          "bytes_per_packet", "bytes_per_sec", "byte_ratio"]
MAX_COL_MISSING = 0.40   # colonne supprimée si > 40 % de vides
MAX_ROW_MISSING = 0.50   # ligne supprimée si > 50 % des features vides
IQR_FACTOR = 3.0         # bornes "extrêmes" calculées sur l'échelle log
TEST_SIZE = 0.20

LOG: list = []
STEP = {"name": ""}


def banner(n, title):
    STEP["name"] = f"{n}. {title}"
    print("\n" + "=" * 74)
    print(f"ÉTAPE {n} : {title}")
    print("=" * 74)


def log(msg):
    print(f"  - {msg}")
    LOG.append((STEP["name"], msg))


def table(df_or_series, indent="      "):
    for line in df_or_series.to_string().splitlines():
        print(indent + line)


# ----------------------------------------------------------------------------
# Pipeline
# ----------------------------------------------------------------------------
def main(inp: str, out: str, seed: int):
    # ---- 1. Chargement ------------------------------------------------------
    banner(1, "Chargement du fichier Excel")
    df = pd.read_excel(inp, sheet_name="raw_data")
    n_start = len(df)
    log(f"Fichier lu : {inp}")
    log(f"Dimensions : {df.shape[0]} lignes x {df.shape[1]} colonnes")
    log(f"Colonnes : {', '.join(df.columns)}")

    # ---- 2. Diagnostic initial ---------------------------------------------
    banner(2, "Diagnostic initial de la qualité des données")
    miss = df.isna().sum()
    miss = miss[miss > 0].sort_values(ascending=False)
    log(f"Cellules vides au total : {int(df.isna().sum().sum())}")
    print("    Colonnes avec valeurs manquantes (nombre | %) :")
    table(pd.DataFrame({"manquantes": miss, "pct": (miss / len(df) * 100).round(1)}))
    log(f"Lignes strictement dupliquées : {int(df.duplicated().sum())}")
    log(f"Valeurs brutes de '{TARGET}' (incohérences de casse visibles) : "
        f"{sorted(df[TARGET].dropna().astype(str).unique().tolist())}")

    # ---- 3. Standardisation du texte ---------------------------------------
    banner(3, "Standardisation des variables catégorielles (espaces, casse)")
    for col in CAT_COLS + [TARGET]:
        s = df[col].astype("string").str.strip().replace("", pd.NA)
        before = df[col].dropna().astype(str).nunique()
        if col == "protocol" or col == "flag":
            s = s.str.upper()
        elif col == "service":
            s = s.str.lower()
        elif col == TARGET:
            key = s.str.lower().str.replace(r"[^a-z0-9]", "", regex=True)
            mapped = key.map(LABEL_MAP)
            unknown = s[mapped.isna() & s.notna()]
            if len(unknown):
                log(f"ATTENTION labels inconnus mis à vide : {unknown.unique().tolist()}")
            s = mapped
        df[col] = s.astype(object).where(s.notna(), np.nan)
        log(f"{col:<9}: {before} variantes -> {df[col].nunique()} valeurs propres "
            f"{sorted(df[col].dropna().unique().tolist())}")

    # ---- 4. Doublons --------------------------------------------------------
    banner(4, "Suppression des doublons")
    n_dup = int(df.duplicated().sum())
    df = df.drop_duplicates().reset_index(drop=True)
    log(f"{n_dup} ligne(s) dupliquée(s) supprimée(s) -> {len(df)} lignes restantes")

    # ---- 5. Valeurs invalides ----------------------------------------------
    banner(5, "Détection des valeurs impossibles (converties en vide)")
    for col in NUM_COLS:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["timestamp"] = pd.to_datetime(df["timestamp"], errors="coerce")
    for col in PORT_COLS:
        bad = (df[col] < 0) | (df[col] > 65535)
        df.loc[bad, col] = np.nan
        log(f"{col}: {int(bad.sum())} port(s) hors de [0, 65535] -> vide")
    for col in NON_NEGATIVE:
        bad = df[col] < 0
        df.loc[bad, col] = np.nan
        log(f"{col}: {int(bad.sum())} valeur(s) négative(s) -> vide")
    log("Les valeurs extrêmes mais plausibles (ex. 800 paquets lors d'un DoS) sont "
        "conservées : ce sont des signaux d'attaque.")

    # ---- 6. Lignes et colonnes inutilisables --------------------------------
    banner(6, "Suppression des lignes / colonnes inutilisables")
    n_before = len(df)
    df = df[df[TARGET].notna()].reset_index(drop=True)
    log(f"{n_before - len(df)} ligne(s) sans label supprimée(s) (inutilisables en supervisé)")
    feat_cols = NUM_COLS + CAT_COLS
    row_miss = df[feat_cols].isna().mean(axis=1)
    n_before = len(df)
    df = df[row_miss <= MAX_ROW_MISSING].reset_index(drop=True)
    log(f"{n_before - len(df)} ligne(s) avec > {MAX_ROW_MISSING:.0%} de features vides supprimée(s)")
    col_miss = df.isna().mean()
    drop_cols = [c for c in col_miss.index
                 if col_miss[c] > MAX_COL_MISSING and c != TARGET and c not in ID_COLS]
    df = df.drop(columns=drop_cols)
    log(f"Colonnes supprimées (> {MAX_COL_MISSING:.0%} de vides) : {drop_cols or 'aucune'}")
    log(f"Lignes restantes : {len(df)}")

    # ---- 7. Suppression des identifiants ------------------------------------
    banner(7, "Extraction de l'heure puis suppression des colonnes d'identification")
    df["hour"] = df["timestamp"].dt.hour.astype(float)
    df = df.drop(columns=[c for c in ID_COLS if c in df.columns])
    log(f"Colonne 'hour' créée depuis timestamp. Colonnes supprimées : {ID_COLS}")
    log("Raison : un identifiant/IP/port éphémère ferait mémoriser le modèle au lieu "
        "de généraliser.")
    num_cols = [c for c in NUM_COLS + ["hour"] if c in df.columns]
    cat_cols = [c for c in CAT_COLS if c in df.columns]

    # ---- 8. Découpage train / test ------------------------------------------
    banner(8, "Découpage Train / Test stratifié (avant toute statistique)")
    stratify = df[TARGET] if df[TARGET].value_counts().min() >= 2 else None
    train, test = train_test_split(df, test_size=TEST_SIZE, random_state=seed, stratify=stratify)
    train, test = train.reset_index(drop=True), test.reset_index(drop=True)
    log(f"Train : {len(train)} lignes | Test : {len(test)} lignes (random_state={seed})")
    print("    Répartition des classes :")
    table(pd.DataFrame({"train": train[TARGET].value_counts(),
                        "test": test[TARGET].value_counts()}).fillna(0).astype(int))

    # ---- 9. Imputation ------------------------------------------------------
    banner(9, "Remplissage des valeurs manquantes (statistiques du TRAIN)")
    impute_values = {}
    for col in num_cols:
        impute_values[col] = float(train[col].median())
    for col in cat_cols:
        impute_values[col] = train[col].mode().iloc[0]
    for col, val in impute_values.items():
        n_tr, n_te = int(train[col].isna().sum()), int(test[col].isna().sum())
        train[col] = train[col].fillna(val)
        test[col] = test[col].fillna(val)
        kind = "médiane" if col in num_cols else "mode"
        log(f"{col:<14}: {n_tr + n_te:>2} remplie(s) avec {kind} = {val}")
    log(f"Vides restants -> train: {int(train.isna().sum().sum())}, test: {int(test.isna().sum().sum())}")

    # ---- 10. Création de nouvelles features ---------------------------------
    banner(10, "Création de features (feature engineering)")
    for d in (train, test):
        d["total_bytes"] = d["src_bytes"] + d["dst_bytes"]
        d["total_packets"] = d["packets_fwd"] + d["packets_bwd"]
        d["bytes_per_packet"] = d["total_bytes"] / d["total_packets"].clip(lower=1)
        d["bytes_per_sec"] = d["total_bytes"] / d["duration"].clip(lower=0.001)
        d["byte_ratio"] = d["src_bytes"] / (d["dst_bytes"] + 1)
        d["is_well_known_port"] = (d["dst_port"] < 1024).astype(int)
        d["has_failed_login"] = (d["failed_logins"] > 0).astype(int)
        d["is_night"] = ((d["hour"] < 7) | (d["hour"] >= 20)).astype(int)
    log("Créées : total_bytes, total_packets, bytes_per_packet, bytes_per_sec, "
        "byte_ratio, is_well_known_port, has_failed_login, is_night")

    # ---- 11. Transformation log + outliers ----------------------------------
    banner(11, "Transformation log1p et plafonnement des valeurs extrêmes")
    for d in (train, test):
        for col in SKEWED:
            d[col] = np.log1p(d[col])
    log(f"log1p appliqué aux {len(SKEWED)} variables très asymétriques (réduit l'écart d'échelle)")
    bounds = {}
    for col in SKEWED:
        q1, q3 = train[col].quantile([.25, .75])
        if q3 - q1 == 0:   # variable quasi constante (ex. failed_logins) : IQR nul -> pas de plafonnement
            log(f"{col}: IQR = 0 (surtout des zéros) -> plafonnement ignoré pour ne pas écraser le signal d'attaque")
            continue
        bounds[col] = (q1 - IQR_FACTOR * (q3 - q1), q3 + IQR_FACTOR * (q3 - q1))
    for name, d in (("train", train), ("test", test)):
        touched = pd.Series(False, index=d.index)
        per_col = {}
        for col, (lo, hi) in bounds.items():
            mask = (d[col] < lo) | (d[col] > hi)
            if mask.any():
                per_col[col] = int(mask.sum())
            d[col] = d[col].clip(lo, hi)
            touched |= mask
        details = ", ".join(f"{c}={n}" for c, n in per_col.items()) or "aucune"
        log(f"{name}: {int(touched.sum())} ligne(s) plafonnée(s) (bornes Q1/Q3 +- {IQR_FACTOR}xIQR "
            f"sur échelle log) -> {details}")
        if touched.any():
            log(f"{name}: classes des lignes plafonnées : "
                f"{d.loc[touched, TARGET].value_counts().to_dict()}")

    # ---- 12. Encodage -------------------------------------------------------
    banner(12, "Encodage des variables catégorielles et de la cible")
    train_oh = pd.get_dummies(train[cat_cols], dtype=int)
    test_oh = pd.get_dummies(test[cat_cols], dtype=int).reindex(columns=train_oh.columns, fill_value=0)
    train = pd.concat([train.drop(columns=cat_cols), train_oh], axis=1)
    test = pd.concat([test.drop(columns=cat_cols), test_oh], axis=1)
    log(f"One-hot de {cat_cols} -> {train_oh.shape[1]} colonnes binaires")
    log("Les colonnes du test sont alignées sur celles du train (catégorie inconnue = que des 0)")
    mapping = {name: i for i, name in enumerate(LABEL_ORDER)}
    for d in (train, test):
        d["is_attack"] = (d[TARGET] != "Normal").astype(int)
        d[TARGET] = d[TARGET].map(mapping).astype(int)
    log(f"Cible '{TARGET}' encodée : {mapping}")
    log("Colonne 'is_attack' ajoutée : 0 = Normal, 1 = attaque (pour une détection binaire)")

    # ---- 13. Normalisation --------------------------------------------------
    banner(13, "Normalisation StandardScaler (moyenne 0, écart-type 1, fit sur le TRAIN)")
    to_scale = SKEWED + [c for c in ["dst_port", "hour", "packets_fwd"] if c not in SKEWED]
    to_scale = list(dict.fromkeys(to_scale))
    scaler = StandardScaler().fit(train[to_scale])
    train[to_scale] = scaler.transform(train[to_scale])
    test[to_scale] = scaler.transform(test[to_scale])
    log(f"{len(to_scale)} colonnes continues normalisées : {to_scale}")
    log("Colonnes binaires / one-hot non normalisées (déjà en 0/1)")
    scaler_df = pd.DataFrame({"feature": to_scale, "mean": scaler.mean_, "scale": scaler.scale_})

    # ---- 14. Contrôles finaux ----------------------------------------------
    banner(14, "Contrôles de qualité finaux")
    feature_cols = [c for c in train.columns if c not in (TARGET, "is_attack")]
    train = train[feature_cols + [TARGET, "is_attack"]].round(6)
    test = test[feature_cols + [TARGET, "is_attack"]].round(6)
    assert train.isna().sum().sum() == 0 and test.isna().sum().sum() == 0, "NaN restants"
    assert np.isfinite(train.to_numpy(dtype=float)).all(), "inf dans train"
    assert np.isfinite(test.to_numpy(dtype=float)).all(), "inf dans test"
    assert list(train.columns) == list(test.columns), "colonnes train/test différentes"
    assert all(pd.api.types.is_numeric_dtype(t) for t in train.dtypes), "colonne non numérique"
    log("OK : aucune valeur manquante, aucune valeur infinie")
    log("OK : toutes les colonnes sont numériques et identiques entre train et test")
    log(f"Train final : {train.shape[0]} x {train.shape[1]} | Test final : {test.shape[0]} x {test.shape[1]}")
    log(f"Features : {len(feature_cols)} | Cible : '{TARGET}' (4 classes) + 'is_attack' (binaire)")
    log(f"Lignes : {n_start} au départ -> {len(train) + len(test)} après nettoyage")

    # ---- 15. Export ---------------------------------------------------------
    banner(15, "Export de l'Excel propre")
    label_df = pd.DataFrame({"label_name": list(mapping), "label_code": list(mapping.values())})
    report_df = pd.DataFrame(LOG, columns=["step", "detail"])
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        train.to_excel(xw, sheet_name="train", index=False)
        test.to_excel(xw, sheet_name="test", index=False)
        label_df.to_excel(xw, sheet_name="label_mapping", index=False)
        scaler_df.to_excel(xw, sheet_name="scaler_params", index=False)
        report_df.to_excel(xw, sheet_name="cleaning_report", index=False)
    style_workbook(out)
    log(f"Fichier créé : {out}")
    print("\nTerminé. Chargement dans Python :")
    print(f"    train = pd.read_excel('{out}', sheet_name='train')")
    print(f"    X, y = train.drop(columns=['{TARGET}', 'is_attack']), train['{TARGET}']\n")


def style_workbook(path: str):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 11), 90)
        ws.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="cyber_data_raw.xlsx")
    ap.add_argument("--output", default="cyber_data_clean.xlsx")
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()
    main(a.input, a.output, a.seed)
