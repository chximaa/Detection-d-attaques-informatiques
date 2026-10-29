"""
generate_dataset.py
-------------------
Génère un jeu de données SIMULÉ de trafic réseau pour un modèle de détection
d'attaques, avec des défauts volontaires (valeurs manquantes, doublons,
incohérences de casse, valeurs impossibles, outliers) afin de tester le
script de nettoyage `clean_dataset.py`.

Sortie : cyber_data_raw.xlsx (200 lignes)

Classes simulées : Normal, DoS, PortScan, BruteForce
Utilisation   : python generate_dataset.py
"""
import argparse
from datetime import datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

N_TOTAL = 200   # nombre de lignes final (doublons inclus)
N_DUP = 5       # lignes dupliquées volontairement

SERVICE_PORT = {"http": 80, "https": 443, "dns": 53, "ssh": 22,
                "ftp": 21, "smtp": 25, "rdp": 3389}

DESCRIPTIONS = {
    "flow_id": "Identifiant du flux réseau",
    "timestamp": "Date et heure du début du flux",
    "src_ip": "Adresse IP source",
    "dst_ip": "Adresse IP destination (serveurs internes)",
    "src_port": "Port source (éphémère)",
    "dst_port": "Port destination (0-65535)",
    "protocol": "Protocole de transport : TCP / UDP / ICMP",
    "service": "Service réseau visé (http, ssh, dns, ...)",
    "flag": "Etat de la connexion (SF, S0, REJ, RSTR)",
    "duration": "Durée du flux en secondes",
    "src_bytes": "Octets envoyés par la source",
    "dst_bytes": "Octets envoyés par la destination",
    "packets_fwd": "Nombre de paquets source -> destination",
    "packets_bwd": "Nombre de paquets destination -> source",
    "failed_logins": "Nombre d'échecs d'authentification",
    "conn_count": "Connexions depuis la même source sur les 2 dernières secondes",
    "label": "Cible : Normal, DoS, PortScan, BruteForce",
}


def build_dataframe(rng: np.random.Generator) -> pd.DataFrame:
    servers = [f"10.0.0.{i}" for i in (10, 11, 12, 20, 30)]
    attackers = {c: [f"{rng.integers(11, 223)}.{rng.integers(0, 255)}."
                     f"{rng.integers(0, 255)}.{rng.integers(1, 254)}" for _ in range(4)]
                 for c in ("DoS", "PortScan", "BruteForce")}

    def stamp(night: bool) -> datetime:
        hour = rng.choice([0, 1, 2, 3, 4, 5, 6, 21, 22, 23]) if night else rng.integers(8, 19)
        return datetime(2026, 9, 1) + timedelta(
            days=int(rng.integers(0, 30)), hours=int(hour),
            minutes=int(rng.integers(0, 60)), seconds=int(rng.integers(0, 60)))

    def normal():
        service = rng.choice(["http", "https", "dns", "ssh", "ftp", "smtp"],
                             p=[.25, .35, .20, .08, .06, .06])
        proto, port = ("UDP" if service == "dns" else "TCP"), SERVICE_PORT[service]
        if rng.random() < 0.05:
            proto, service, port = "ICMP", "other", 0
        return dict(
            timestamp=stamp(False),
            src_ip=f"192.168.{rng.integers(0, 5)}.{rng.integers(2, 250)}",
            dst_ip=rng.choice(servers), src_port=int(rng.integers(1024, 65535)),
            dst_port=port, protocol=proto, service=service,
            flag=rng.choice(["SF", "REJ", "RSTR"], p=[.85, .05, .10]),
            duration=round(float(rng.lognormal(0.0, 1.0)), 3),
            src_bytes=int(rng.lognormal(np.log(800), 0.9)),
            dst_bytes=int(rng.lognormal(np.log(4000), 1.1)),
            packets_fwd=int(rng.poisson(15)) + 1, packets_bwd=int(rng.poisson(18)) + 1,
            failed_logins=int(rng.random() < 0.05), conn_count=int(rng.poisson(3)) + 1,
            label="Normal")

    def dos():
        service = rng.choice(["http", "https", "other"], p=[.45, .35, .20])
        return dict(
            timestamp=stamp(True), src_ip=rng.choice(attackers["DoS"]),
            dst_ip=rng.choice(servers[:3]), src_port=int(rng.integers(1024, 65535)),
            dst_port=SERVICE_PORT.get(service, 80),
            protocol=rng.choice(["TCP", "UDP", "ICMP"], p=[.6, .25, .15]), service=service,
            flag=rng.choice(["S0", "SF"], p=[.8, .2]),
            duration=round(float(rng.uniform(0.001, 0.3)), 3),
            src_bytes=int(rng.lognormal(np.log(600), 0.6)), dst_bytes=int(rng.integers(0, 200)),
            packets_fwd=int(rng.poisson(800)) + 50, packets_bwd=int(rng.poisson(2)),
            failed_logins=0, conn_count=int(rng.poisson(150)) + 50, label="DoS")

    def portscan():
        return dict(
            timestamp=stamp(True), src_ip=rng.choice(attackers["PortScan"]),
            dst_ip=rng.choice(servers), src_port=int(rng.integers(1024, 65535)),
            dst_port=int(rng.integers(1, 65535)), protocol="TCP", service="other",
            flag=rng.choice(["REJ", "S0", "RSTR"], p=[.4, .4, .2]),
            duration=round(float(rng.uniform(0.0005, 0.05)), 4),
            src_bytes=int(rng.integers(40, 80)), dst_bytes=int(rng.integers(0, 60)),
            packets_fwd=int(rng.integers(1, 4)), packets_bwd=int(rng.integers(0, 2)),
            failed_logins=0, conn_count=int(rng.poisson(60)) + 20, label="PortScan")

    def brute():
        service = rng.choice(["ssh", "ftp", "rdp"], p=[.5, .3, .2])
        return dict(
            timestamp=stamp(True), src_ip=rng.choice(attackers["BruteForce"]),
            dst_ip=rng.choice(servers[:3]), src_port=int(rng.integers(1024, 65535)),
            dst_port=SERVICE_PORT[service], protocol="TCP", service=service,
            flag=rng.choice(["SF", "RSTR"], p=[.5, .5]),
            duration=round(float(rng.uniform(1, 12)), 3),
            src_bytes=int(rng.lognormal(np.log(900), 0.4)),
            dst_bytes=int(rng.lognormal(np.log(700), 0.4)),
            packets_fwd=int(rng.poisson(40)) + 10, packets_bwd=int(rng.poisson(35)) + 10,
            failed_logins=int(rng.poisson(12)) + 4, conn_count=int(rng.poisson(20)) + 5,
            label="BruteForce")

    n_unique = N_TOTAL - N_DUP
    counts = {"Normal": int(n_unique * .56), "DoS": int(n_unique * .15),
              "PortScan": int(n_unique * .15)}
    counts["BruteForce"] = n_unique - sum(counts.values())
    makers = {"Normal": normal, "DoS": dos, "PortScan": portscan, "BruteForce": brute}
    rows = [makers[c]() for c, k in counts.items() for _ in range(k)]

    df = pd.DataFrame(rows).sort_values("timestamp").reset_index(drop=True)
    df.insert(0, "flow_id", [f"F{i:04d}" for i in range(1, len(df) + 1)])
    return df


def inject_defects(df: pd.DataFrame, rng: np.random.Generator) -> pd.DataFrame:
    df = df.copy()
    n = len(df)

    # 1) Valeurs manquantes (cellules vides) : taux différent par colonne
    miss_rates = {"src_ip": .02, "src_port": .05, "dst_port": .04, "protocol": .05,
                  "service": .06, "flag": .05, "duration": .07, "src_bytes": .06,
                  "dst_bytes": .06, "packets_fwd": .05, "packets_bwd": .05,
                  "failed_logins": .04, "conn_count": .06}
    for col, rate in miss_rates.items():
        df[col] = df[col].astype(object)
        df.loc[rng.random(n) < rate, col] = np.nan
    df.loc[rng.choice(n, 3, replace=False), "label"] = np.nan  # 3 labels manquants

    # 2) Incohérences de casse / espaces
    df["protocol"] = [v if pd.isna(v) or rng.random() > .15 else
                      rng.choice([v.lower(), f" {v} ", v.capitalize()]) for v in df["protocol"]]
    df["service"] = [v if pd.isna(v) or rng.random() > .10 else v.upper() for v in df["service"]]
    variants = {"Normal": ["normal", "NORMAL"], "DoS": ["dos", "DOS ", "Dos"],
                "PortScan": ["portscan", "Port_Scan", "PORTSCAN"],
                "BruteForce": ["brute_force", "bruteforce", "Brute Force"]}
    df["label"] = [v if pd.isna(v) or rng.random() > .25 else rng.choice(variants[v])
                   for v in df["label"]]

    # 3) Valeurs impossibles (sur des cellules non vides)
    def pick(col, k):
        return rng.choice(df.index[df[col].notna()], k, replace=False)
    df.loc[pick("duration", 3), "duration"] = -1.0                 # durée négative
    df.loc[pick("src_port", 1), "src_port"] = 70000                # port > 65535
    df.loc[pick("dst_port", 1), "dst_port"] = -1                   # port négatif
    df.loc[pick("packets_fwd", 1), "packets_fwd"] = -5             # compteur négatif
    df.loc[pick("src_bytes", 3), "src_bytes"] = 200_000_000        # outliers extrêmes

    # 4) Doublons exacts, insérés à des positions aléatoires
    dup = df.iloc[rng.choice(n, N_DUP, replace=False)]
    df = pd.concat([df, dup], ignore_index=True)
    return df.iloc[rng.permutation(len(df))].reset_index(drop=True)


def save_excel(df: pd.DataFrame, path: str) -> None:
    dico = pd.DataFrame({"column": list(DESCRIPTIONS), "description": list(DESCRIPTIONS.values())})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="raw_data", index=False)
        dico.to_excel(xw, sheet_name="data_dictionary", index=False)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 70)
        ws.freeze_panes = "A2"
    wb["raw_data"].column_dimensions["B"].width = 20
    for r in range(2, wb["raw_data"].max_row + 1):
        wb["raw_data"].cell(r, 2).number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", default="cyber_data_raw.xlsx")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    rng = np.random.default_rng(args.seed)
    data = inject_defects(build_dataframe(rng), rng)
    save_excel(data, args.output)

    print(f"Fichier généré : {args.output}  ({len(data)} lignes, {data.shape[1]} colonnes)")
    print(f"Cellules vides : {int(data.isna().sum().sum())}")
    print(f"Lignes dupliquées : {int(data.duplicated().sum())}")
